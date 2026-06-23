import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import datetime

# Cores
COLOR_DARK_BLUE = "0B5394"
COLOR_LIGHT_BLUE = "9FC5E8"
COLOR_VERY_LIGHT_BLUE = "CFE2F3"
COLOR_WHITE = "FFFFFF"

# Fontes
FONT_TITLE = Font(name="Arial", size=11, bold=True, color=COLOR_WHITE)
FONT_HEADER = Font(name="Arial", size=10, bold=True)
FONT_DATA = Font(name="Arial", size=10)
FONT_DATA_BOLD = Font(name="Arial", size=10, bold=True)

# Preenchimentos
FILL_DARK_BLUE = PatternFill(start_color=COLOR_DARK_BLUE, end_color=COLOR_DARK_BLUE, fill_type="solid")
FILL_LIGHT_BLUE = PatternFill(start_color=COLOR_LIGHT_BLUE, end_color=COLOR_LIGHT_BLUE, fill_type="solid")
FILL_VERY_LIGHT_BLUE = PatternFill(start_color=COLOR_VERY_LIGHT_BLUE, end_color=COLOR_VERY_LIGHT_BLUE, fill_type="solid")

# Bordas
BORDER_THIN = Border(
    left=Side(style='thin', color='C0C0C0'),
    right=Side(style='thin', color='C0C0C0'),
    top=Side(style='thin', color='C0C0C0'),
    bottom=Side(style='thin', color='C0C0C0')
)

def build_sheet_formalizacao(ws, entidades):
    ws.title = "FORMALIZAÇÃO"
    
    # Linha 1: 3 banners agrupados com merge
    # SITUAÇÃO (A1:B1)
    ws.merge_cells("A1:B1")
    ws["A1"] = "SITUAÇÃO"
    # ENTIDADE (C1:D1)
    ws.merge_cells("C1:D1")
    ws["C1"] = "ENTIDADE"
    # EMENDA (E1:L1)
    ws.merge_cells("E1:L1")
    ws["E1"] = "EMENDA"
    
    # Estilizar linha 1
    for col in range(1, 13):
        cell = ws.cell(row=1, column=col)
        cell.fill = FILL_DARK_BLUE
        cell.font = FONT_TITLE
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER_THIN
    ws.row_dimensions[1].height = 25
    
    # Linha 2: Subcabeçalhos
    headers = [
        "STATUS", "HISTÓRICO", "NOME", "CNPJ", "PA Emenda",
        "Localização do PA Emenda", "Emenda Alterada?", "PA Formalização",
        "N.º", "Vereador", "Justificativa", "Valor"
    ]
    for idx, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=idx, value=header)
        cell.fill = FILL_LIGHT_BLUE
        cell.font = FONT_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_THIN
    ws.row_dimensions[2].height = 25
    
    # Linhas 3+: Dados das entidades
    for row_idx, ent in enumerate(entidades, 3):
        # Fundo alternado
        fill = FILL_VERY_LIGHT_BLUE if row_idx % 2 == 1 else PatternFill(fill_type=None)
        
        # Valor formatado
        val = ent.get("valor")
        if val is not None and val != "":
            try:
                val = float(val)
            except:
                val = 0.0
        else:
            val = 0.0
            
        row_data = [
            ent.get("situacao") or "",
            ent.get("historico") or "",
            ent.get("razao_social") or "",
            ent.get("cnpj") or "",
            ent.get("pa_emenda") or "",
            ent.get("localizacao_pa_emenda") or "",
            ent.get("emenda_alterada") or "",
            ent.get("pa_formalizacao") or "",
            ent.get("numero_emenda") or "",
            ent.get("vereador") or "",
            ent.get("justificativa") or "",
            val
        ]
        
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = FONT_DATA
            cell.border = BORDER_THIN
            if fill.fill_type:
                cell.fill = fill
                
            # Alinhamentos específicos
            if col_idx == 2:  # Histórico
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            elif col_idx == 12:  # Valor
                cell.number_format = '"R$ "#,##0.00;"R$ "(#,##0.00);"R$ "-";@'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx in (1, 4, 5, 7, 8, 9):  # Status, CNPJ, PAs, N.º, Alterada
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
    # Configurar larguras fixas
    widths = {
        'A': 15, 'B': 55, 'C': 45, 'D': 20, 'E': 15, 'F': 15,
        'G': 15, 'H': 15, 'I': 8, 'J': 20, 'K': 35, 'L': 15
    }
    for col_letter, w in widths.items():
        ws.column_dimensions[col_letter].width = w

def build_sheet_parceria(ws, entidades):
    ws.title = "DADOS DA PARCERIA"
    
    last_col = 43
    last_col_letter = get_column_letter(last_col)
    
    # Banner A1:AQ1 (AQ é coluna 43)
    ws.merge_cells(f"A1:{last_col_letter}1")
    ws["A1"] = "RELAÇÃO DE AJUSTES VIGENTES COM TERCEIRO SETOR - SAÚDE"
    
    # Estilizar linha 1
    for col in range(1, last_col + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = FILL_DARK_BLUE
        cell.font = FONT_TITLE
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER_THIN
    ws.row_dimensions[1].height = 25
    
    # Cabeçalhos na linha 2 (altura ≈ 130)
    ws.row_dimensions[2].height = 130
    
    categorias = ["Saúde Mental", "Fisioterapia", "Fono", "Animal", "Outros"]
    especialidades = [
        "Academia Clínica", "Acupuntura", "Assistente Social", "Atividade Educativa", "Educador Físico",
        "Fisio", "Fono", "Hidroginástica/\nHidroterapia", "Massoterapeuta", "Médico (Neurologista)",
        "Musicoterapia", "Neuropediatra", "Neuropsicologia", "Nutricionista", "Odonto", "Oficinas Lúdicas",
        "Oftalmologia", "Ortopedista", "Pediatria", "Pilates", "Psicologia", "Psicanalista", "Psiquiatria",
        "Psicomotricista", "Psicopedagogo", "Práticas Integrativas", "Reflexologia", "T.O.", "Veterinário"
    ]
    
    headers = (
        ["ENTIDADE", "AJUSTE", "INÍCIO DAS ATIVIDADES", "TÉRMINO DAS ATIVIDADES", "GESTOR DA PARCERIA", "PROJETO"]
        + categorias
        + ["ATENDIMENTO", "META (Mês)"]
        + especialidades
        + ["RESPONSÁVEL PELA ENTIDADE"]
    )
    
    for idx, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=idx, value=header)
        cell.fill = FILL_LIGHT_BLUE
        cell.font = FONT_HEADER
        cell.border = BORDER_THIN
        
        # Rotação para categorias e especialidades
        is_rotated = (idx in range(7, 12)) or (idx in range(14, 43))
        if is_rotated:
            cell.alignment = Alignment(text_rotation=90, horizontal="center", vertical="center", wrap_text=True)
        else:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            
    # Preencher dados das entidades (Linhas 3+)
    for row_idx, ent in enumerate(entidades, 3):
        fill = FILL_VERY_LIGHT_BLUE if row_idx % 2 == 1 else PatternFill(fill_type=None)
        
        def fmt_date(d):
            if isinstance(d, (datetime.date, datetime.datetime)):
                return d.strftime('%d/%m/%Y')
            elif isinstance(d, str) and d:
                # Caso venha YYYY-MM-DD
                parts = d.split('-')
                if len(parts) == 3 and len(parts[0]) == 4:
                    return f"{parts[2]}/{parts[1]}/{parts[0]}"
                return d
            return ""
            
        inicio = fmt_date(ent.get("inicio_atividades"))
        termino = fmt_date(ent.get("termino_atividades"))
        
        cats_db = ent.get("categorias") or {}
        if isinstance(cats_db, str):
            import json
            try:
                cats_db = json.loads(cats_db)
            except:
                cats_db = {}
        
        cats_values = []
        for cat in categorias:
            if isinstance(cats_db, dict):
                cats_values.append("X" if cats_db.get(cat) else "")
            elif isinstance(cats_db, list):
                cats_values.append("X" if cat in cats_db else "")
            else:
                cats_values.append("")
            
        esp_db = ent.get("especialidades") or {}
        if isinstance(esp_db, str):
            import json
            try:
                esp_db = json.loads(esp_db)
            except:
                esp_db = {}
        if not isinstance(esp_db, dict):
            esp_db = {}
            
        esp_values = []
        for esp in especialidades:
            val_esp = esp_db.get(esp)
            if val_esp is None:
                found = False
                for k, v in esp_db.items():
                    if k.replace("\n", " ").replace("\r", "").strip() == esp.replace("\n", " ").replace("\r", "").strip():
                        val_esp = v
                        found = True
                        break
                if not found:
                    val_esp = ""
            esp_values.append(val_esp if val_esp is not None else "")
            
        row_data = (
            [
                ent.get("razao_social") or "",
                ent.get("ajuste_termo") or "",
                inicio,
                termino,
                ent.get("gestor_parceria") or "",
                ent.get("projeto") or ""
            ]
            + cats_values
            + [
                ent.get("atendimento_descricao") or "",
                ent.get("meta_mes_atendimentos") or ""
            ]
            + esp_values
            + [ent.get("responsavel_entidade") or ""]
        )
        
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = FONT_DATA
            cell.border = BORDER_THIN
            if fill.fill_type:
                cell.fill = fill
                
            is_center = (col_idx in range(3, 5)) or (col_idx in range(7, 12)) or (col_idx in range(14, 43))
            if is_center:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                
    # Larguras das colunas
    for col in range(1, last_col + 1):
        col_letter = get_column_letter(col)
        if col in (1, 5, 6, 12):
            ws.column_dimensions[col_letter].width = 30
        elif col in (2, 13, 43):
            ws.column_dimensions[col_letter].width = 20
        elif col in (3, 4):
            ws.column_dimensions[col_letter].width = 15
        elif (col in range(7, 12)) or (col in range(14, 43)):
            ws.column_dimensions[col_letter].width = 6

def build_sheet_repasse(ws, entidade, repasses):
    rs = entidade.get("razao_social") or "REPASSE"
    for c in ['\\', '/', '?', '*', '[', ']', ':']:
        rs = rs.replace(c, '')
    base_title = "Repasse e Prestação de contas"
    
    sheet_title = base_title
    counter = 1
    if hasattr(ws, 'parent') and ws.parent:
        while sheet_title in ws.parent.sheetnames and ws.parent[sheet_title] != ws:
            sheet_title = f"{base_title[:22]}_{counter}"
            counter += 1
            
    ws.title = sheet_title
    
    inicio_dt = entidade.get("inicio_atividades")
    ano_base = None
    if isinstance(inicio_dt, (datetime.date, datetime.datetime)):
        ano_base = inicio_dt.year
    elif isinstance(inicio_dt, str) and inicio_dt:
        try:
            ano_base = int(inicio_dt.split("-")[0])
        except:
            pass
            
    if not ano_base:
        ano_base = datetime.datetime.now().year
        
    ano_anterior = str(ano_base - 1)
    meses_colunas = [ano_anterior] + [f"{str(m).zfill(2)}.{ano_base}" for m in range(1, 13)]
    
    # Mesclar A1:D1
    ws.merge_cells("A1:D1")
    ws["A1"] = "ENTIDADE"
    
    ws.row_dimensions[1].height = 25
    for col_idx in range(1, 5 + len(meses_colunas)):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = FILL_DARK_BLUE
        cell.font = FONT_TITLE
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER_THIN
        
    for idx, mes in enumerate(meses_colunas, 5):
        ws.cell(row=1, column=idx, value=mes)
        
    # Estrutura bloco esquerdo
    ws.cell(row=2, column=1, value="ENTIDADE:")
    ws.cell(row=3, column=1, value="CNPJ")
    ws.cell(row=4, column=1, value="CÓD. SCIM:")
    ws.cell(row=5, column=1, value="AJUSTE:")
    ws.cell(row=6, column=1, value="P.A. EMPENHO:")
    ws.cell(row=7, column=1, value="OBJETO:")
    ws.merge_cells("A7:A21")
    
    ws.cell(row=2, column=2, value=entidade.get("razao_social") or "")
    ws.cell(row=3, column=2, value=entidade.get("cnpj") or "")
    ws.cell(row=4, column=2, value=entidade.get("cod_scim") or "")
    ws.cell(row=5, column=2, value=entidade.get("ajuste_termo") or "")
    ws.cell(row=6, column=2, value=entidade.get("pa_empenho") or "")
    ws.cell(row=7, column=2, value=entidade.get("objeto_descricao") or "")
    ws.merge_cells("B7:B21")
    
    for r in range(2, 22):
        for c in (1, 2):
            cell = ws.cell(row=r, column=c)
            cell.font = FONT_DATA_BOLD if c == 1 else FONT_DATA
            cell.border = BORDER_THIN
            if c == 1:
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                
    # Seções coluna C
    ws.cell(row=2, column=3, value="REPASSE")
    ws.merge_cells("C2:C9")
    ws.cell(row=10, column=3, value="PRESTAÇÃO DE CONTAS")
    ws.merge_cells("C10:C15")
    ws.cell(row=16, column=3, value="MTS")
    ws.merge_cells("C16:C21")
    
    rotulos = {
        2: "OFÍCIO",
        3: "PERÍODO",
        4: "PARCELA ($)",
        5: "RETENÇÃO ($)",
        6: "VALOR FINAL",
        7: "VENCIMENTO",
        8: "P.A.",
        9: "DATA - PAGAMENTO",
        10: "OFÍCIO",
        11: "DATA - ENTREGA",
        12: "P.A.",
        13: "SUGESTÃO DE GLOSA ($)",
        14: "RECONSIDERAÇÃO ($)",
        15: "SUGESTÃO DE GLOSA (P.A.)",
        16: "Cadastro (Termo)",
        17: "Atualização / Aditamento",
        18: "Lançado Valores? (TSS)",
        19: "Prestado Contas? (ENTIDADE)",
        20: "Análise da PC? (COMISSÃO)",
        21: "Parecer Final?"
    }
    
    for r, rotulo in rotulos.items():
        ws.cell(row=r, column=4, value=rotulo)
        
    for r in range(2, 22):
        for c in (3, 4):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER_THIN
            cell.fill = FILL_LIGHT_BLUE
            cell.font = FONT_HEADER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            
    repasses_map = {}
    for rep in repasses:
        mes_ref = rep.get("mes_referencia")
        if mes_ref:
            repasses_map[str(mes_ref).strip()] = rep
            
    fields_mapping = [
        (2, "repasse_oficio", "str"),
        (3, "repasse_periodo", "str"),
        (4, "repasse_parcela", "currency"),
        (5, "repasse_retencao", "currency"),
        (6, "repasse_valor_final", "currency"),
        (7, "repasse_vencimento", "date"),
        (8, "repasse_pa", "str"),
        (9, "repasse_data_pagamento", "date"),
        (10, "prestacao_oficio", "str"),
        (11, "prestacao_data_entrega", "date"),
        (12, "prestacao_pa", "str"),
        (13, "prestacao_sugestao_glosa", "currency"),
        (14, "prestacao_reconsideracao", "currency"),
        (15, "prestacao_sugestao_glosa_pa", "str"),
    ]
    
    def format_dt(val):
        if isinstance(val, (datetime.date, datetime.datetime)):
            return val.strftime('%d/%m/%Y')
        elif isinstance(val, str) and val:
            parts = val.split('-')
            if len(parts) == 3 and len(parts[0]) == 4:
                return f"{parts[2]}/{parts[1]}/{parts[0]}"
            return val
        return ""

    for col_idx, mes in enumerate(meses_colunas, 5):
        rep = repasses_map.get(mes)
        
        for row_num, field_name, field_type in fields_mapping:
            cell = ws.cell(row=row_num, column=col_idx)
            cell.border = BORDER_THIN
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
            if rep is None:
                cell.value = "------"
                continue
                
            if field_name == "prestacao_sugestao_glosa_pa":
                cell.value = ""
                continue
                
            val = rep.get(field_name)
            
            if field_type == "currency":
                if val is not None and val != "":
                    try:
                        val_num = float(val)
                        cell.value = val_num
                        cell.number_format = '"R$ "#,##0.00;"R$ "(#,##0.00);"R$ "-";@'
                    except:
                        cell.value = "------"
                else:
                    cell.value = 0.0
                    cell.number_format = '"R$ "#,##0.00;"R$ "(#,##0.00);"R$ "-";@'
            elif field_type == "date":
                formatted_d = format_dt(val)
                cell.value = formatted_d if formatted_d else "------"
            else:
                cell.value = str(val) if val is not None and str(val).strip() != "" else "------"
                
        for row_num in range(16, 22):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.border = BORDER_THIN
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.value = "" if rep is not None else "------"
            
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 50
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 25
    for col_idx in range(5, 5 + len(meses_colunas)):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = 15
